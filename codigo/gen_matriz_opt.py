"""
gen_matriz_opt.py
================================================================================
Genera la Matriz de Respuesta del Autor de la VERSION DE OPTIMIZACION (articulo/main.tex):
  entregas/Matriz_Respuesta_Optimizacion.docx
Las ubicaciones (seccion, ecuacion, tabla, figura y pagina) se leen de articulo/main.aux y las cifras de
articulo/numeros.tex, de modo que la matriz siempre coincide con el manuscrito compilado.
Requiere python-docx (pip install python-docx) y haber compilado articulo/main.tex.
Uso:  python gen_matriz_opt.py
"""
import _raiz  # noqa: F401  (cwd = carpeta del proyecto)
import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = "entregas/Matriz_Respuesta_Optimizacion.docx"
TITULO = "Matriz de Respuesta del Autor — mejora del controlador difuso"
SUBTITULO = ("Mejora de un controlador de semáforos de lógica difusa para una intersección aislada: optimización del controlador de "
             "Pappis y Mamdani (1977)")
AUTORES = "Raul Nolberto Quispe Saya y Kevin Luis Valero Ccacya — Escuela Profesional de Ingeniería Electrónica, UNSA"


# ---------------------------------------------------------------- lectura de etiquetas y cifras
def leer_etiquetas(ruta="articulo/main.aux"):
    L = {}
    for m in re.finditer(r"\\newlabel\{([^}]+)\}\{\{(?:\\mbox\s*\{)?([^}]+?)\}?\}\{(\d+)\}", open(ruta, encoding="utf-8").read()):
        L[m.group(1)] = (m.group(2), int(m.group(3)))
    return L


def leer_numeros(ruta="articulo/numeros.tex"):
    return {m.group(1): m.group(2).replace("$-$", "−") for m in re.finditer(r"\\newcommand\{\\(\w+)\}\{([^}]*)\}", open(ruta, encoding="utf-8").read())}


L = leer_etiquetas()
LS = leer_etiquetas("articulo/suplemento.aux")
N = leer_numeros()


def refs(clave):
    """Ubicacion en el suplemento (suplemento.pdf), leida de suplemento.aux."""
    num, pag = LS[clave]
    return f"Tabla {num} del suplemento (p. {pag})" if clave.startswith("tab:") else f"Sec. {num} del suplemento (p. {pag})"


def ref(clave, pagina=True):
    num, pag = L[clave]
    if clave.startswith("sec:"):
        txt = f"Sec. {num}"
    elif clave.startswith("eq:"):
        txt = f"Ec. ({num})"
    elif clave.startswith("tab:"):
        txt = f"Tabla {num}"
    else:
        txt = f"Fig. {num}"
    return f"{txt} (p. {pag})" if pagina else txt


def paginas_main():
    import subprocess
    out = subprocess.run(["pdfinfo", "articulo/main.pdf"], capture_output=True, text=True).stdout
    return re.search(r"Pages:\s+(\d+)", out).group(1)


NP = paginas_main()

INTRO = [
    "Esta matriz acompaña al manuscrito (articulo/main.pdf, "
    f"{NP} páginas con referencias, IEEEtran, y su suplemento articulo/suplemento.pdf). El manuscrito mejora el controlador de semáforos con lógica difusa de "
    "Pappis y Mamdani (1977) mediante optimización con un protocolo reproducible. Parte de la réplica de ese controlador hecha antes por los autores "
    "(articulo_replica/ y rama replica del repositorio), que sirve de línea base y de validación del simulador.",
    "La Parte A indica qué ocurre con cada una de las 17 observaciones de la revisión anterior, hechas sobre la réplica, en este manuscrito. La Parte B responde a los puntos que la réplica no "
    "podía atender por su propio alcance (novedad, comparadores, información usada, generalización). Las ubicaciones se leen del manuscrito compilado; las cifras salen "
    "de resultados/opt/evaluacion_final.json mediante codigo/gen_tablas_opt.py. Todo se reproduce con python codigo/ejecutar_pruebas.py.",
]

# (N.º, titulo, categoria, estado, respuesta y ubicacion)
PARTE_A = [
    ("1", "Operador «más que» (mt)", "Crítica", "Conservada",
     f"La definición correcta (0 para x<x0 y 1−μA(x) para x≥x0; lt∪mt = complemento y lt∩mt = ∅ en todo el universo, incluido x0) se conserva en {ref('eq:mtlt')}, {ref('sec:estructura')}; "
     "es lo que ejecuta el código (mt_arr, lt_queue) y lo comprueban test_paper_tables.py y test_parametros.py."),
    ("2", "Parámetros latentes de CALIBRADA", "Crítica", "Conservada",
     f"La lectura calibrada se declara de forma explícita (a_e = min(e, Q_G+Σx_i), q_e = Q_V, q_mode='current', predicted_q=False) y como inferencia por falta de datos en {ref('sec:lecturas')}. "
     "Los parámetros se guardan resueltos en SimConfig y CALIBRADA; ninguno se ignora (test_parametros.py)."),
    ("3", "Scripts ausentes y LaTeX obsoleto", "Crítica", "Atendida",
     "Todo archivo que el manuscrito o el suplemento citan existe: test_archivos_citados.py (réplica) y test_archivos_citados_opt.py (esta versión) lo comprueban y forman parte de "
     "ejecutar_pruebas.py. Las tablas y las cifras del texto no se escriben a mano: gen_tablas_opt.py las genera desde los resultados (numeros.tex, tabla_*.tex)."),
    ("4", "Retardo residual al final del horizonte", "Importante", "Superada y ampliada",
     f"El retardo del manuscrito es ahora siempre el que incluye la espera residual, {ref('eq:dm')} (D = suma por segundo de las colas); la identidad es una deducción y se verifica en "
     "test_sim_batch.py (D = D_ns + D_ew y equivalencia con el simulador de referencia). La convención del original (solo vehículos servidos) se reporta aparte (delay_servidos en evaluacion_final.json)."),
    ("5", "Línea base redondeada y alcance de la comparación", "Importante", "Atendida por completo",
     f"Además del ciclo fijo de Webster se comparan un actuado por detector, un actuado afinado con el mismo presupuesto, la presión máxima, la cota óptima por programación dinámica, "
     f"una política neuronal y FLC-O ({ref('tab:controladores')}, {ref('tab:conjuntos')}). El actuado afinado es un rival más fuerte que el original: "
     f"FLC-O lo mejora en {N['ParMediaActAf']} % por par, con {N['ParMejorActAf']} pares significativamente mejores y {N['ParPeorActAf']} peores ({ref('tab:comparaciones')})."),
    ("6", "Números aleatorios comunes, dispersión y pruebas no paramétricas", "Importante", "Atendida",
     f"Conjunto oficial con 20 réplicas y números aleatorios comunes; media, desviación estándar muestral (n−1) e IC 95 % por par; Wilcoxon pareado con Holm "
     f"({ref('tab:oficial')}, {ref('tab:comparaciones')}; todos los pares en la {refs('tab:spares')}). Los conjuntos adicionales usan 64 escenarios."),
    ("7", "Activación de reglas con una semilla y 1800 s", "Importante", "Superada por el cambio de alcance",
     f"El manuscrito ya no reporta activación de reglas del controlador de 1977 con una semilla. La interpretación de FLC-O se da con su base de reglas aprendida ({ref('tab:reglasflc')}, {ref('fig:difuso')}) "
     "y todos los resultados usan 20 réplicas de 7200 s o 64 escenarios de 7200 s. El análisis agregado de disparos de la réplica sigue en el artículo de réplica."),
    ("8", "Literatura desactualizada y celda incompleta", "Importante", "Atendida",
     f"{ref('sec:sota')} cubre hasta 2026 e incorpora las referencias de optimización (CMA-ES, presión máxima, programación dinámica, ley de Little, Takagi–Sugeno). "
     "Los DOI se comprueban contra Crossref con codigo/verificar_referencias.py; la celda incompleta de la tabla de antecedentes ya no existe (la tabla de 180 combinaciones queda en el artículo de réplica)."),
    ("9", "Duplicación estructural", "Importante", "Atendida",
     f"El manuscrito tiene {NP} páginas (la versión de réplica, 15); cada ecuación se define una sola vez ({ref('sec:modelo', False)} y {ref('sec:pappis', False)}) y los detalles de la réplica se remiten al artículo de réplica, no se repiten."),
    ("10", "Entorno de software y hardware y complejidad", "Menor", "Atendida",
     f"Sección «Disponibilidad del código» (p. {L['sec:disp'][1]}) y {refs('sec:s-repro')}: Python 3.13, NumPy 2.5, generador PCG64 con semillas fijas, 4 procesos. Costo: con un núcleo, unos 10 ms por escenario de 7200 s con el simulador vectorizado "
     "frente a unos 230 ms con el de referencia. Cada decisión de FLC-O es O(27) por segundo."),
    ("11", "Desviación estándar de población y de muestra", "Menor", "Atendida",
     "Toda desviación estándar es la muestral (n−1): estadistica.ic95 y las tablas del manuscrito y del suplemento (test_estadistica.py)."),
    ("12", "2.0 % frente a 2.1 % en la Tabla I", "Menor", "Conservada",
     f"El error de 2.1 % de la réplica se cita igual en {ref('sec:lecturas')} y en el resumen; coincide con el artículo de réplica."),
    ("13", "Ecuación suelta y numeración", "Menor", "Atendida",
     f"Las {len([k for k in L if k.startswith('eq:')])} ecuaciones están numeradas en orden y todas se citan en el texto."),
    ("14", "Colisión de notación del símbolo q", "Menor", "Atendida",
     f"Tasa de llegada q̄_i y colas Q; símbolos sin repetirse y con tabla de notación ({ref('tab:notacion')})."),
    ("15", "Hoja de ruta incompleta", "Menor", "Atendida",
     f"El último párrafo de la Introducción enumera todas las secciones ({ref('sec:intro', False)})."),
    ("16", "Errores tipográficos", "Menor", "Atendida", "Revisión del texto y compilación con pdfLaTeX y BibTeX sin errores ni referencias sin definir."),
    ("17", "Gramática y separadores decimales", "Menor", "Atendida", "Separador decimal con punto en todo el manuscrito (es-nodecimaldot) y revisión gramatical."),
]

PARTE_B = [
    ("B1", "Novedad «incremental» (replicación y auditoría)",
     f"Se acepta que la réplica por sí sola era incremental. El manuscrito añade una contribución nueva: FLC-O, un difuso de 27 reglas con 42 parámetros entrenados con CMA-ES, junto con un simulador "
     f"vectorizado exacto, una cota óptima y un protocolo de prueba separado de entrenamiento ({ref('sec:intro', False)}, contribuciones 1 a 5).",
     f"Retardo medio oficial: {N['DelayOfPap']} s/veh (Pappis y Mamdani) → {N['DelayOfFlc']} (FLC-O), −{N['RedFlc']} %; significativamente mejor en {N['ParMejorPap']} de 16 pares y peor en {N['ParPeorPap']}."),
    ("B2", "Comparadores modernos ausentes",
     f"Se agregan: actuado afinado, presión máxima, cota óptima por programación dinámica y una política neuronal entrenada con la misma optimización; todos con el protocolo de la Sección VII ({ref('tab:controladores')}).",
     f"Oficial: Webster {N['DelayOfFijo']}, actuado afinado {N['DelayOfActAf']}, presión máxima {N['DelayOfMax']}, óptimo {N['DelayOfMdp']}, política neuronal {N['DelayOfMlp']}, FLC-O {N['DelayOfFlc']} s/veh."),
    ("B3", "Variante de Takagi y Sugeno no implementada",
     f"FLC-O es un sistema de Mamdani con consecuentes de un solo punto, equivalente a uno de Takagi y Sugeno de orden cero ({ref('eq:salida')}, {ref('sec:flc')}); es la variante que se entrena y se evalúa.",
     f"{ref('tab:reglasflc')} y {ref('fig:difuso')} muestran sus 27 reglas y sus funciones de pertenencia aprendidas."),
    ("B4", "Información que usa el controlador (vista adelante de 10 s)",
     f"Se declara que la vista adelante es un supuesto del original y se mide cuánto de la mejora depende de ella ({ref('sec:res-abl')}). Sin ella la mejora casi desaparece: el difuso causal no supera al actuado afinado y la mejor variante causal lo supera solo en {N['RedActMlpC']} %.",
     f"Sin vista adelante: FLC-O causal {N['DelayOfFlcC']} ({N['RedActFlcC']} % frente al actuado afinado) y política neuronal causal {N['DelayOfMlpC']} s/veh ({N['RedActMlpC']} %), frente a {N['DelayOfActAf']} del actuado afinado; con vista adelante, FLC-O da {N['DelayOfFlc']} ({N['RedActAfOf']} %)."),
    ("B5", "Generalización más allá de la Tabla VIII",
     f"Cuatro conjuntos de prueba que no se usaron para entrenar: oficial, 64 caudales no vistos, demanda variable y cerca de la saturación ({ref('tab:conjuntos')}).",
     f"Reducción frente a Pappis y Mamdani: oficial {N['RedPapOf']} %, no vistos {N['RedPapNv']} %, variable {N['RedPapVa']} %, saturado {N['RedPapSa']} %."),
    ("B6", "Sobreajuste y fuga de información",
     f"Entrenamiento con caudales sorteados (semillas 100000+), validación con otros (200000+) y prueba en la Tabla VIII; el verde mínimo se eligió con validación, no con la prueba. La sensibilidad a la semilla de entrenamiento (tres semillas) está en la {refs('tab:ssemillas')}.",
     f"Hiperparámetros: {refs('sec:s-entrena')}; semillas: {refs('sec:s-semillas')}."),
    ("B7", "Limitaciones declaradas",
     f"{ref('sec:discusion')}: la mejora depende de la vista adelante; cerca de la saturación el resultado no es el mejor; todo es simulación con un cruce y llegadas de Bernoulli; el retardo corregido no es comparable con la Tabla VIII del original.",
     f"Saturado: FLC-O {N['DelaySaFlc']}, Pappis y Mamdani {N['DelaySaPap']}, actuado afinado {N['DelaySaActAf']} s/veh."),
]

NOTAS = [
    "El repositorio público https://github.com/INVESTVALQUISP/PROYECTO_INVESTIGACION (rama main) contiene el código, las pruebas, los resultados (resultados/opt/) y los manuscritos (articulo/ para la mejora y articulo_replica/ para la réplica de partida); la rama replica conserva el estado anterior del repositorio, centrado solo en la réplica.",
    "Puntos abiertos que dependen de los autores: validar con datos de campo, y entrenar con un rango de caudales que cubra la saturación (trabajo futuro del manuscrito).",
]


# ---------------------------------------------------------------- formato del documento
def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def escribir(celda, texto, negrita=False, tam=7.5, color=None):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(texto)
    r.font.size = Pt(tam)
    r.bold = negrita
    if color:
        r.font.color.rgb = RGBColor.from_string(color)


def repetir_encabezado(fila):
    trPr = fila._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def fijar_anchos(tabla, anchos):
    tblPr = tabla._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for i, col in enumerate(tabla.columns):
        col.width = anchos[i]
    for fila in tabla.rows:
        for i, c in enumerate(fila.cells):
            c.width = anchos[i]


def titulo(doc, texto, tam=10):
    p = doc.add_paragraph()
    r = p.add_run(texto)
    r.bold = True
    r.font.size = Pt(tam)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(9)
    p = doc.add_paragraph()
    r = p.add_run(TITULO)
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    r = p.add_run(SUBTITULO)
    r.italic = True
    r.font.size = Pt(10)
    p = doc.add_paragraph()
    r = p.add_run(AUTORES)
    r.font.size = Pt(9)
    for t in INTRO:
        p = doc.add_paragraph(t)
        p.paragraph_format.space_after = Pt(4)
        for r in p.runs:
            r.font.size = Pt(9)

    titulo(doc, "Parte A. Las 17 observaciones de la revisión anterior en este manuscrito")
    anchos = [Cm(1.8), Cm(5.2), Cm(3.0), Cm(17.3)]
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Table Grid"
    tb.autofit = False
    for i, tx in enumerate(["Obs.", "Observación", "Estado", "Respuesta, ubicación y evidencia en este manuscrito"]):
        escribir(tb.rows[0].cells[i], tx, True, 8, "FFFFFF")
        sombrear(tb.rows[0].cells[i], "1F3864")
    repetir_encabezado(tb.rows[0])
    colores = {"Crítica": "FDE9E7", "Importante": "FFF4D6", "Menor": "E8F1E4"}
    for n, tit, cat, est, resp in PARTE_A:
        c = tb.add_row().cells
        escribir(c[0], f"{n}\n{cat}", True, 8)
        sombrear(c[0], colores[cat])
        escribir(c[1], tit, True, 7.5)
        escribir(c[2], est, False, 7.5)
        escribir(c[3], resp, False, 7.5)
    fijar_anchos(tb, anchos)
    doc.add_paragraph()

    titulo(doc, "Parte B. Puntos que la réplica no podía atender por su alcance")
    anchos = [Cm(1.4), Cm(4.6), Cm(12.2), Cm(9.1)]
    tb = doc.add_table(rows=1, cols=4)
    tb.style = "Table Grid"
    tb.autofit = False
    for i, tx in enumerate(["Punto", "Observación", "Respuesta y ubicación", "Evidencia"]):
        escribir(tb.rows[0].cells[i], tx, True, 8, "FFFFFF")
        sombrear(tb.rows[0].cells[i], "1F3864")
    repetir_encabezado(tb.rows[0])
    for n, tit, resp, evi in PARTE_B:
        c = tb.add_row().cells
        escribir(c[0], n, True, 8)
        escribir(c[1], tit, True, 7.5)
        escribir(c[2], resp, False, 7.5)
        escribir(c[3], evi, False, 7.5)
    fijar_anchos(tb, anchos)
    doc.add_paragraph()

    titulo(doc, "Notas finales")
    for tx in NOTAS:
        q = doc.add_paragraph(tx)
        for r in q.runs:
            r.font.size = Pt(8.5)

    cp = doc.core_properties
    cp.author = "Raul Nolberto Quispe Saya; Kevin Luis Valero Ccacya"
    cp.last_modified_by = cp.author
    cp.title = TITULO
    cp.comments = ""
    os.makedirs("entregas", exist_ok=True)
    doc.save(OUT)
    print("guardado", OUT)


if __name__ == "__main__":
    main()
